"""Genera la plantilla de Google Sheets (xlsx) de la réplica de Loftus y Burns (1982).
Uso:  python3 crear_plantilla_xlsx.py [salida.xlsx] [--demo]
  --demo  rellena 60 filas simuladas (solo para comprobar las fórmulas)."""
import json, sys, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

H = json.load(open(sys.argv[3] if len(sys.argv) > 3 else "columnas.json"))
out = sys.argv[1] if len(sys.argv) > 1 else "Loftus_Burns_datos.xlsx"
demo = "--demo" in sys.argv
AZUL, MARINO, NARANJA = "0047E8", "001447", "FF9606"
F = "Arial"
wb = Workbook()

def cab(ws, fila, textos, color=AZUL):
    for i, t in enumerate(textos, 1):
        c = ws.cell(fila, i, t)
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)

# ---------------- Respuestas ----------------
ws = wb.active; ws.title = "Respuestas"
cab(ws, 1, H); ws.freeze_panes = "D2"
for i, h in enumerate(H, 1):
    ws.column_dimensions[get_column_letter(i)].width = max(9, min(26, len(h) + 3))
col = {h: get_column_letter(i + 1) for i, h in enumerate(H)}
N = 3000
def R(h): return f"Respuestas!${col[h]}$2:${col[h]}${N}"

if demo:
    random.seed(4)
    for r in range(2, 62):
        g = "experimental" if r % 2 else "control"; e = g == "experimental"
        row = {h: "" for h in H}
        row.update(id=f"SIM{r}", cohorte="2026-27", condicion=g, cambios_pestana=1 if r == 7 else 0,
                   conocia="Sí" if r == 9 else "No", condiciones="No" if r == 12 else "Sí")
        ac = 0
        for k in range(1, 11):
            ok = int(random.random() < (0.66 if e else 0.8)); row[f"p{k:02d}_ok"] = ok; ac += ok
        row["aciertos_relleno"] = ac
        row["c1_recuerdo_ok"] = int(random.random() < (0.1 if e else 0.35))
        row["c2_reconocimiento_ok"] = int(random.random() < (0.3 if e else 0.55))
        row["malestar"] = random.choice([3, 4, 4, 5]) if e else random.choice([1, 1, 2, 2, 3])
        for i, h in enumerate(H, 1): ws.cell(r, i, row[h])

# ---------------- Resumen (fórmulas en vivo) ----------------
rs = wb.create_sheet("Resumen", 0)
rs.column_dimensions["A"].width = 46
for c in "BCDE": rs.column_dimensions[c].width = 18
rs["A1"] = "Réplica de Loftus y Burns (1982) · Resumen en vivo"
rs["A1"].font = Font(name=F, size=16, bold=True, color=AZUL)
rs["A2"] = "ESIC University · Métodos, Diseño y Técnicas de Investigación en Psicología · Tema 4"
rs["A2"].font = Font(name=F, size=10, color=MARINO)
rs["A4"] = "Cohorte a analizar (vacío = todas):"; rs["B4"] = ""
rs["B4"].fill = PatternFill("solid", fgColor="FFF3E0"); rs["B4"].font = Font(name=F, bold=True)
rs["A5"] = "Excluidos automáticamente: conocían el estudio, no vieron bien las imágenes o cambiaron de pestaña."
rs["A5"].font = Font(name=F, size=9, italic=True, color="43506B")

def mask(g):
    return (f'({R("condicion")}="{g}")*({R("conocia")}<>"Sí")*({R("condiciones")}<>"No")'
            f'*({R("cambios_pestana")}=0)*((({R("cohorte")}=$B$4)+($B$4=""))>0)')
ME, MC = mask("experimental"), mask("control")

cab(rs, 7, ["Aciertos de relleno (VD, 0-10)", "Experimental", "Control", "Diferencia"])
x = R("aciertos_relleno")
filas = [
    ("n incluidos", f"=SUMPRODUCT({ME})", f"=SUMPRODUCT({MC})", ""),
    ("Media", f"=IFERROR(SUMPRODUCT({ME}*{x})/B8,\"\")", f"=IFERROR(SUMPRODUCT({MC}*{x})/C8,\"\")", "=IFERROR(B9-C9,\"\")"),
    ("Desviación típica", f"=IFERROR(SQRT(SUMPRODUCT({ME}*({x}-B9)^2)/(B8-1)),\"\")",
     f"=IFERROR(SQRT(SUMPRODUCT({MC}*({x}-C9)^2)/(C8-1)),\"\")", ""),
    ("% de aciertos", "=IFERROR(B9/10,\"\")", "=IFERROR(C9/10,\"\")", ""),
]
for i, f in enumerate(filas):
    for j, v in enumerate(f): rs.cell(8 + i, 1 + j, v)
for c in ("B11", "C11"): rs[c].number_format = "0.0%"
for c in ("B9", "C9", "D9", "B10", "C10"): rs[c].number_format = "0.00"

cab(rs, 13, ["t de Welch (medias independientes)", "Valor", "", ""])
tests = [
    ("t", "=IFERROR(D9/SQRT(B10^2/B8+C10^2/C8),\"\")"),
    ("grados de libertad (Welch)", "=IFERROR((B10^2/B8+C10^2/C8)^2/((B10^2/B8)^2/(B8-1)+(C10^2/C8)^2/(C8-1)),\"\")"),
    ("p bilateral (TDIST trunca los gl)", "=IFERROR(TDIST(ABS(B14),B15,2),\"\")"),
    ("d de Cohen", "=IFERROR(D9/SQRT(((B8-1)*B10^2+(C8-1)*C10^2)/(B8+C8-2)),\"\")"),
]
for i, (a, v) in enumerate(tests):
    rs.cell(14 + i, 1, a); rs.cell(14 + i, 2, v).number_format = "0.000"
rs["A18"] = "La U de Mann-Whitney (no paramétrica) se calcula en resultados.html."
rs["A18"].font = Font(name=F, size=9, italic=True, color="43506B")

def critico(fila0, titulo, campo, azar=None):
    cab(rs, fila0, [titulo, "Experimental", "Control", ""])
    ok = R(campo); r = fila0 + 1
    rs.cell(r, 1, "Aciertos (n)"); rs.cell(r, 2, f"=SUMPRODUCT({ME}*{ok})"); rs.cell(r, 3, f"=SUMPRODUCT({MC}*{ok})")
    rs.cell(r + 1, 1, "% de aciertos")
    rs.cell(r + 1, 2, f"=IFERROR(B{r}/B8,\"\")").number_format = "0.0%"
    rs.cell(r + 1, 3, f"=IFERROR(C{r}/C8,\"\")").number_format = "0.0%"
    a, c = f"B{r}", f"C{r}"; b, d = f"(B8-B{r})", f"(C8-C{r})"; n = "(B8+C8)"
    rs.cell(r + 2, 1, "χ² (1 gl, corrección de Yates)")
    rs.cell(r + 2, 2, f"=IFERROR({n}*MAX(0,ABS({a}*{d}-{b}*{c})-{n}/2)^2/(({a}+{b})*({c}+{d})*({a}+{c})*({b}+{d})),\"\")").number_format = "0.00"
    rs.cell(r + 3, 1, "p"); rs.cell(r + 3, 2, f"=IFERROR(CHIDIST(B{r+2},1),\"\")").number_format = "0.000"
    if azar: rs.cell(r + 1, 4, f"azar = {azar}")
critico(20, "Ítem crítico: recuerdo del 17", "c1_recuerdo_ok")
critico(26, "Ítem crítico: reconocimiento (10/13/1/17)", "c2_reconocimiento_ok", "25 %")

cab(rs, 32, ["Control de la manipulación: malestar (1-5)", "Experimental", "Control", ""])
m = R("malestar")
rs["A33"] = "Media"; rs["B33"] = f"=IFERROR(SUMPRODUCT({ME}*{m})/B8,\"\")"; rs["C33"] = f"=IFERROR(SUMPRODUCT({MC}*{m})/C8,\"\")"
rs["A34"] = "Desviación típica"
rs["B34"] = f"=IFERROR(SQRT(SUMPRODUCT({ME}*({m}-B33)^2)/(B8-1)),\"\")"
rs["C34"] = f"=IFERROR(SQRT(SUMPRODUCT({MC}*({m}-C33)^2)/(C8-1)),\"\")"
rs["A35"] = "t de Welch"; rs["B35"] = "=IFERROR((B33-C33)/SQRT(B34^2/B8+C34^2/C8),\"\")"
rs["A36"] = "p bilateral"
rs["B36"] = "=IFERROR(TDIST(ABS(B35),(B34^2/B8+C34^2/C8)^2/((B34^2/B8)^2/(B8-1)+(C34^2/C8)^2/(C8-1)),2),\"\")"
for c in ("B33", "C33", "B34", "C34", "B35"): rs[c].number_format = "0.00"
rs["B36"].number_format = "0.000"
rs["A38"] = "Referencia original (Exp. 1): recuerdo del 17 = 4,3 % (violento) vs. 27,9 % (no violento); χ²(1) = 21,72."
rs["A38"].font = Font(name=F, size=9, italic=True, color="43506B")
for row in rs.iter_rows(min_row=7, max_row=36):
    for c in row:
        if c.font.color is None or c.font.color.rgb != "00FFFFFF": c.font = Font(name=F, bold=c.font.bold)

# ---------------- Clave ----------------
ck = wb.create_sheet("Clave")
cab(ck, 1, ["Ítem", "Pregunta", "Respuesta correcta", "Tipo", "Dónde se ve"])
for w, c in zip([18, 70, 32, 12, 38], "ABCDE"): ck.column_dimensions[c].width = w
for i, it in enumerate(json.load(open("preguntas.json")), 2):
    for j, v in enumerate([it["id"], it["texto"], str(it["correcta"]), "crítico" if it.get("critica") else "relleno", it.get("fuente", "")], 1):
        ck.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Asignaciones ----------------
asg = wb.create_sheet("Asignaciones"); cab(asg, 1, ["fecha", "cohorte", "condicion"], MARINO)
for c in "ABC": asg.column_dimensions[c].width = 22

# ---------------- Leeme ----------------
lm = wb.create_sheet("Leeme"); lm.column_dimensions["A"].width = 28; lm.column_dimensions["B"].width = 90
cab(lm, 1, ["Pestaña / columna", "Descripción"], MARINO)
texto = [
    ("Resumen", "Estadísticos en vivo calculados con fórmulas a partir de 'Respuestas'. Cambia B4 para filtrar una cohorte."),
    ("Respuestas", "Una fila por participante. La escribe la web (no editar las cabeceras de la fila 1)."),
    ("Clave", "Respuestas correctas del cuestionario (solo informativa; la codificación la hace la web)."),
    ("Asignaciones", "Registro de la aleatorización equilibrada (una fila por consentimiento aceptado)."),
    ("condicion", "experimental (final impactante) / control (final neutro)"),
    ("pXX / pXX_ok / pXX_conf", "respuesta bruta / 1 = acierto, 0 = error / seguridad 1-5"),
    ("aciertos_relleno", "VD principal: suma de aciertos en los 10 ítems de relleno"),
    ("c1_recuerdo_ok", "ítem crítico: recuerdo libre del número de la camiseta (17)"),
    ("c2_reconocimiento_ok", "ítem crítico: reconocimiento entre 10, 13, 1 y 17"),
    ("malestar", "control de la manipulación (1-5)"),
    ("cambios_pestana, conocia, condiciones", "criterios de exclusión"),
    ("stroop_*", "rendimiento en la tarea distractora"),
    ("Apps Script", "Extensiones > Apps Script: pegar apps-script/Code.gs e implementar como aplicación web (ver README)."),
]
for i, (a, b) in enumerate(texto, 2):
    lm.cell(i, 1, a).font = Font(name=F, bold=True); lm.cell(i, 2, b).alignment = Alignment(wrap_text=True)
wb.save(out); print("ok", out)
